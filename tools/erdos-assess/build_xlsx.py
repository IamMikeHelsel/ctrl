"""Build the workbook. Usage: build_xlsx.py <open.csv> <solved_aligned.json> <assess_dir> <jsp_all.json> <out.xlsx>"""
import csv, json, glob, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
ocsv, sal, adir, jspf, out = sys.argv[1:6]
J = {j['id']: j for j in json.load(open(jspf))}
S = {e['number']: e for e in json.load(open(sal))}
A = [a for f in sorted(glob.glob(f'{adir}/solved_*.json')) for a in json.load(open(f))]
def usd(p):
    if not p or p == 'no': return 0
    m = re.search(r'[\d,]+', p); v = int(m.group(0).replace(',', '')) if m else 0
    return round(v * 1.3) if '£' in p else round(v / 85) if '₹' in p else v
solved = []
for a in A:
    s = S[str(a['number'])]; ja = s.get('jsp_aligned')
    solved.append({'erdos_no': int(a['number']), 'prize': s['prize'] if s['prize'] != 'no' else '', 'prize_usd_approx': usd(s['prize']),
        'priority_1to5': int(a['priority']), 'formalization_effort_1to5': int(a['formalization_effort']),
        'status': s['status'], 'summary': a['summary'], 'solution_source': a['solution_source'], 'mathlib_gaps': a['mathlib_gaps'],
        'note_flags': a.get('note_flags', ''), 'lean_statement_in_formal_conjectures': s['lean_statement'],
        'jsp_id_unconfirmed': ja['id'] if ja else '', 'jsp_title_unconfirmed': ja['title'] if ja else '', 'jsp_status': ja['status'] if ja else '',
        'tags': '; '.join(s['tags']), 'erdos_url': f"https://www.erdosproblems.com/{a['number']}",
        'formal_conjectures_url': f"https://github.com/google-deepmind/formal-conjectures/blob/main/FormalConjectures/ErdosProblems/{a['number']}.lean" if s['lean_statement'] == 'yes' else '',
        'jsp_url': f"https://github.com/TheJustinSunPrize/awards/blob/main/problems/{J[ja['id']]['file']}#{ja['id']}" if ja else ''})
solved.sort(key=lambda r: (-r['priority_1to5'], r['formalization_effort_1to5'], -r['prize_usd_approx'], r['erdos_no']))
opened = list(csv.DictReader(open(ocsv)))
for r in opened:
    for k in ('erdos_no', 'prize_usd_approx', 'difficulty_1to5', 'ai_lean_fit_1to5', 'tractability_score', 'comments'): r[k] = int(r[k])
wb = Workbook(); ws = wb.active; ws.title = 'README'
readme = [
 ('Erdős problems — assessment workbook', ''), ('Generated', '2026-10-17 by Claude Code; source data teorth/erdosproblems @ b850d1a (2026-10-16), erdosproblems.com pages, TheJustinSunPrize/awards catalog.'),
 ('Tab: Open', '642 problems the database lists as open (incl. falsifiable/decidable/verifiable/ZFC-limited). Sorted by tractability_score = (6 - difficulty) + ai_lean_fit.'),
 ('Tab: Formalizer targets', '263 problems solved informally but with no Lean solution yet. Being first to formalize is what the Justin Sun Prize has paid formalizer-only awards for.'),
 ('difficulty_1to5', '1 = plausibly closable soon … 5 = famous/deep. LLM judgement from site notes; not verified.'),
 ('ai_lean_fit_1to5', '5 = crisp/finite/elementary with a Lean statement ready; 1 = heavy analysis absent from Mathlib.'),
 ('possibly_resolved', 'Notes mention a recent claimed or formalised (dis)proof — check before investing. 48 flagged.'),
 ('jsp_*', 'Justin Sun Prize catalog match. Open tab: agent-confirmed with confidence. Formalizer tab: automatic order alignment, UNCONFIRMED.'),
 ('prize_usd_approx', 'Erdős prize in USD (£×1.3, ₹÷85). Paid by the Combinatorics Foundation (Steve Butler) only after journal publication + proof Erdős offered it.'),
 ('Caveats', 'All ratings are quick model assessments from scraped notes, not literature reviews. Statement texts are not reproduced (site owner asks that its text not be reused); follow erdos_url.'),
]
for row in readme: ws.append(row)
ws['A1'].font = Font(bold=True, size=14); ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 140
for c in ws['B']: c.alignment = Alignment(wrap_text=True, vertical='top')
for name, rows in (('Open', opened), ('Formalizer targets', solved)):
    w = wb.create_sheet(name); cols = list(rows[0]); w.append(cols)
    for r in rows: w.append([r[c] for c in cols])
    for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496')
    w.freeze_panes = 'B2'; w.auto_filter.ref = w.dimensions
    for i, c in enumerate(cols, 1):
        w.column_dimensions[get_column_letter(i)].width = 60 if c in ('summary', 'stepping_stone', 'resolved_note', 'jsp_title_unconfirmed', 'mathlib_gaps', 'note_flags', 'solution_source') else 14
wb.save(out); print('wrote', out, len(opened), len(solved))
